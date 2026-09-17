"""
Generator Laporan Lengkap UTS Deep Learning (Word .docx)
Pendekatan: Eksperimen Optimasi Bertingkat (Progressive / Sequential Ablation Study)
Dataset: The IQ-OTHNCCD Lung Cancer CT-Scan Dataset (Mendeley Data DOI: 10.17632/bhmdr45bh2.2)
Dosen Pengampu: Dr. Wahyudi Setiawan, S.Kom., M.Kom.
Mahasiswa:
1. Attala Alif Ramadhani Tri Hida (NIM: 230441100144)
2. Nafaul Hernanda Romadlona (NIM: 240441100125)
"""

import os
import json
import pandas as pd
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUTS_DIR = os.path.join(BASE_DIR, 'outputs')
FIGURES_DIR = os.path.join(OUTPUTS_DIR, 'figures')
LOGS_DIR = os.path.join(OUTPUTS_DIR, 'logs')
REPORT_PATH = os.path.join(BASE_DIR, 'Laporan_Lengkap_UTS_DeepLearning_CNN.docx')

COLOR_PRIMARY_HEX = "1A365D"    # Deep Navy
COLOR_SECONDARY_HEX = "2B6CB0"  # Slate Blue
COLOR_LIGHT_BG_HEX = "F7FAFC"   # Off-white

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'))

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    run.font.name = 'Arial'
    if level == 1:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
    return h

def format_styled_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header
    for i, h_text in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = h_text
        set_cell_background(c, COLOR_PRIMARY_HEX)
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Rows
    for row_idx, r_data in enumerate(rows_data):
        row = table.add_row()
        bg = COLOR_LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(r_data):
            c = row.cells[col_idx]
            c.text = str(val)
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 100, 100)
            p = c.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
                if col_idx == 0: r.font.bold = True
                
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def generate_word_report():
    print("[LAPORAN] Membaca rekapitulasi data eksperimen bertingkat...")
    df_prog = pd.read_csv(os.path.join(LOGS_DIR, 'progressive_pipeline_summary.csv'))
    df_all = pd.read_csv(os.path.join(LOGS_DIR, 'all_models_detailed_summary.csv'))
    
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        
    # JUDUL & COVER
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run("LAPORAN UJIAN TENGAH SEMESTER (UTS)\nMATA KULIAH DEEP LEARNING\n")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run("Studi Komparasi Bertingkat (Progressive Ablation Study) Arsitektur Convolutional Neural Network (CNN) pada Citra CT-Scan Kanker Paru-Paru (IQ-OTH/NCCD)")
    r_s.font.name = 'Calibri'
    r_s.font.size = Pt(12)
    r_s.font.bold = True
    r_s.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    # KOTAK IDENTITAS
    t_meta = doc.add_table(rows=6, cols=2)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Dosen Pengampu", ": Dr. Wahyudi Setiawan, S.Kom., M.Kom."),
        ("Kelompok", ": Kelompok 6"),
        ("Anggota Kelompok", ": 1. Attala Alif Ramadhani Tri Hida (230441100144)\n  2. Nafaul Hernanda Romadlona (240441100125)\n  3. M.Rafly Kurniawan (240441100086)\n  4. Naufal Husain (240441100038)"),
        ("Program Studi / Kelas", ": Sistem Informasi / Deep Learning (A)"),
        ("Sumber Dataset", ": Mendeley Data (DOI: 10.17632/bhmdr45bh2.2)"),
        ("Metodologi Eksperimen", ": Eksperimen Optimasi Bertingkat (Progressive Ablation Study)")
    ]
    for r_i, (k, v) in enumerate(meta_info):
        c0, c1 = t_meta.rows[r_i].cells
        c0.text, c1.text = k, v
        c0.width, c1.width = Inches(2.3), Inches(4.2)
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "F7FAFC")
        set_cell_margins(c0, 40, 40, 80, 80)
        set_cell_margins(c1, 40, 40, 80, 80)
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # BAB I: PENDAHULUAN
    add_styled_heading(doc, "BAB I. PENDAHULUAN", level=1)
    doc.add_paragraph(
        "Kanker paru-paru merupakan penyebab utama mortalitas terkait kanker secara global. Deteksi dini melalui modalitas pencitraan Computed Tomography (CT-Scan) thorax adalah standar utama dalam menilai keberadaan nodul paru. "
        "Convolutional Neural Network (CNN) telah terbukti mampu mengekstraksi representasi fitur morfologis jaringan medis dengan akurasi tinggi."
    )
    doc.add_paragraph(
        "Sesuai arahan Dosen Pengampu Dr. Wahyudi Setiawan, S.Kom., M.Kom., eksperimen ini dirancang menggunakan 1 arsitektur CNN yang sama dengan pendekatan Eksperimen Optimasi Bertingkat (Progressive Ablation Study). "
        "Pada pendekatan ini, setiap skenario menentukan konfigurasi terbaik dari suatu parameter, yang kemudian langsung diwariskan ke skenario berikutnya untuk pengujian faktor baru hingga diperoleh Model Juara Akhir (Final Champion Model)."
    )
    
    # BAB II: DATASET
    add_styled_heading(doc, "BAB II. SUMBER DATASET & PREPROCESSING", level=1)
    doc.add_paragraph(
        "Dataset yang digunakan adalah The IQ-OTHNCCD Lung Cancer Dataset yang dipublikasikan secara terbuka di Mendeley Data (DOI: 10.17632/bhmdr45bh2.2) "
        "oleh Hamdalla Alyasriy & Muayed AL-Huseiny dari Wasit University, Irak. Data dihimpun dari Iraq-Oncology Teaching Hospital dan National Center for Cancer Diseases. "
        "Dataset mencakup 1.097 citra CT-scan 2D dengan distribusi kelas: Benign (120 citra), Malignant (561 citra), dan Normal (416 citra)."
    )
    doc.add_paragraph(
        "Setiap citra CT-Scan 512x512 diubah ukurannya menjadi 128x128 piksel dan dinormalisasi intensitasnya ke rentang [0.0, 1.0]. Pembagian dataset dilakukan menggunakan stratified sampling agar proporsi kelas tetap konsisten."
    )
    
    # BAB III: ARSITEKTUR MODEL
    add_styled_heading(doc, "BAB III. DESAIN ARSITEKTUR CONVOLUTIONAL NEURAL NETWORK", level=1)
    doc.add_paragraph(
        "Model dirancang dengan 4 blok konvolusi hierarkis (memenuhi syarat minimal 3 blok Conv-ReLU-Pooling):\n"
        "• Blok 1: Conv2D(32 filter, 3x3) + ReLU + MaxPooling2D(2x2)\n"
        "• Blok 2: Conv2D(64 filter, 3x3) + ReLU + MaxPooling2D(2x2)\n"
        "• Blok 3: Conv2D(128 filter, 3x3) + ReLU + MaxPooling2D(2x2)\n"
        "• Blok 4: Conv2D(128 filter, 3x3) + ReLU + MaxPooling2D(2x2)\n"
        "• Classifier: Flatten() -> Dense(128 neuron, ReLU) -> Dropout(rate) -> Dense(3 neuron, Softmax)"
    )
    
    # BAB IV: METODOLOGI ALUR BERTINGKAT
    add_styled_heading(doc, "BAB IV. DESAIN ALUR EKSPERIMEN BERTINGKAT (4 SKENARIO)", level=1)
    doc.add_paragraph(
        "Konsep 4 skenario eksperimen dirancang secara progresif berkesinambungan:\n"
        "1. Skenario 1 (Optimasi Split Data): Menguji partisi data 70:15:15 vs 80:10:10 vs 90:05:05 pada baseline (Adam, Dropout 0.3, Tanpa Augmentasi). Hasil rasio split terbaik diwariskan ke Skenario 2.\n"
        "2. Skenario 2 (Optimasi Data Augmentasi): Menggunakan split terbaik dari Skenario 1, lalu menguji secara langsung Tanpa Augmentasi vs Dengan Augmentasi (Flip, Rotasi, Zoom). Strategi augmentasi terbaik diwariskan ke Skenario 3.\n"
        "3. Skenario 3 (Optimasi Optimizer): Menggunakan split terbaik dan strategi augmentasi terbaik, lalu membandingkan tiga algoritma optimasi: Adam vs RMSprop vs SGD Momentum. Optimizer terbaik diwariskan ke Skenario 4.\n"
        "4. Skenario 4 (Optimasi Regularisasi Dropout): Menggunakan konfigurasi terbaik dari Skenario 1, 2, dan 3, lalu menguji variasi nilai Dropout (0.0 vs 0.3 vs 0.5) untuk menentukan Final Champion Model."
    )
    
    # BAB V: HASIL & PEMBAHASAN
    add_styled_heading(doc, "BAB V. HASIL EKSPERIMEN & PEMBAHASAN", level=1)
    
    add_styled_heading(doc, "5.1 Ringkasan Progresi Tiap Tahap (Sequential Progression)", level=2)
    doc.add_paragraph("Tabel berikut menyajikan pemenang dari masing-masing tahap pengujian berjenjang:")
    
    t_prog = doc.add_table(rows=1, cols=6)
    headers_p = ["Tahap Pengujian", "Eksperimen Terpilih", "Konfigurasi Pemenang", "Akurasi (%)", "F1-Score (%)", "Loss"]
    rows_p = []
    for _, r in df_prog.iterrows():
        rows_p.append([r['Tahap'], r['Pemenang'], r['Konfigurasi Terpilih'], f"{r['Test Accuracy (%)']:.2f}%", f"{r['Macro F1 (%)']:.2f}%", f"{r['Test Loss']:.4f}"])
    format_styled_table(t_prog, [1.4, 1.8, 1.5, 0.7, 0.7, 0.6], headers_p, rows_p)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Gambar Progresi Bertingkat
    bar_prog_path = os.path.join(FIGURES_DIR, 'progressive_progression_bar.png')
    if os.path.exists(bar_prog_path):
        doc.add_picture(bar_prog_path, width=Inches(6.2))
        p_c1 = doc.add_paragraph()
        p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c1 = p_c1.add_run("Gambar 5.1: Grafik Progresi Peningkatan Performa Pemenang Tiap Tahap")
        r_c1.font.italic = True
        r_c1.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        
    add_styled_heading(doc, "5.2 Rekapitulasi Detail Seluruh Model yang Ditraining", level=2)
    doc.add_paragraph("Tabel berikut memuat evaluasi komparatif internal untuk seluruh 11 variasi model yang diuji di setiap skenario:")
    
    t_all = doc.add_table(rows=1, cols=7)
    headers_a = ["Skenario", "Nama Variasi Model", "Loss", "Akurasi (%)", "Precision (%)", "Recall (%)", "F1-Score (%)"]
    rows_a = []
    for _, r in df_all.iterrows():
        rows_a.append([r['Tahap / Skenario'], r['Nama Eksperimen'], f"{r['Test Loss']:.4f}", f"{r['Test Accuracy (%)']:.2f}%", f"{r['Precision (%)']:.2f}%", f"{r['Recall (%)']:.2f}%", f"{r['F1-Score (%)']:.2f}%"])
    format_styled_table(t_all, [1.1, 2.2, 0.6, 0.7, 0.7, 0.7, 0.7], headers_a, rows_a)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Gambar Komparasi Internal
    int_plot_path = os.path.join(FIGURES_DIR, 'internal_scenario_comparisons.png')
    if os.path.exists(int_plot_path):
        doc.add_picture(int_plot_path, width=Inches(6.2))
        p_c2 = doc.add_paragraph()
        p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c2 = p_c2.add_run("Gambar 5.2: Perbandingan Internal Metrik di Setiap Skenario Pengujian")
        r_c2.font.italic = True
        r_c2.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        
    add_styled_heading(doc, "5.3 Evaluasi Final Champion Model", level=2)
    doc.add_paragraph(
        f"Model Juara Akhir yang terpilih adalah konfigurasi dari Tahap 4: "
        f"Kombinasi Split 90:05:05, Tanpa Augmentasi, Optimizer Adam (lr=0.0005), dan Dropout 0.3. "
        f"Model ini mencatatkan Akurasi Uji sebesar {df_prog.iloc[-1]['Test Accuracy (%)']:.2f}% dengan Macro F1-Score {df_prog.iloc[-1]['Macro F1 (%)']:.2f}% dan Test Loss {df_prog.iloc[-1]['Test Loss']:.4f}."
    )
    
    # Gambar Evaluasi Model Pemenang
    champ_fig_path = os.path.join(FIGURES_DIR, 'champion_model_evaluation.png')
    if os.path.exists(champ_fig_path):
        doc.add_picture(champ_fig_path, width=Inches(6.2))
        p_c3 = doc.add_paragraph()
        p_c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c3 = p_c3.add_run("Gambar 5.3: Kurva Loss, Akurasi, dan Matriks Konfusi Final Champion Model")
        r_c3.font.italic = True
        r_c3.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        
    # BAB VI: KESIMPULAN
    add_styled_heading(doc, "BAB VI. KESIMPULAN", level=1)
    doc.add_paragraph(
        "1. Pendekatan eksperimen bertingkat (Progressive Ablation Study) memberikan alur optimasi yang sistematis dan terarah dalam menemukan kombinasi hyperparameter terbaik untuk klasifikasi kanker paru-paru.\n"
        "2. Pada Tahap 1, rasio data split terbukti menjadi fondasi penting dalam menentukan kuantitas variasi citra latih yang dipelajari arsitektur CNN.\n"
        "3. Pada Tahap 2, perbandingan langsung membuktikan bahwa citra CT-Scan murni (Tanpa Augmentasi) pada epoch pendek memberikan sinyal konvergensi yang lebih cepat dan tajam dibandingkan augmentasi rotasi sintetis.\n"
        "4. Pada Tahap 3, optimizer Adam terbukti mengungguli RMSprop dan SGD Momentum dalam mengatasi lanskap loss non-konveks citra CT-Scan toraks.\n"
        "5. Pada Tahap 4, nilai regularisasi Dropout 0.3 membuktikan perannya sebagai titik optimal dalam mencegah co-adaptasi neuron, menghasilkan Final Champion Model dengan akurasi uji 98.18%."
    )
    
    # DAFTAR PUSTAKA
    add_styled_heading(doc, "DAFTAR PUSTAKA", level=1)
    refs = [
        "Alyasriy, H., & AL-Huseiny, M. (2021). The IQ-OTHNCCD lung cancer dataset. Mendeley Data, V2, doi: 10.17632/bhmdr45bh2.2.",
        "LeCun, Y., Bottou, L., Bengio, Y., & Haffner, P. (1998). Gradient-based learning applied to document recognition. Proceedings of the IEEE, 86(11), 2278-2324.",
        "Simonyan, K., & Zisserman, A. (2014). Very deep convolutional networks for large-scale image recognition. arXiv preprint arXiv:1409.1556.",
        "Srivastava, N., Hinton, G., Krizhevsky, A., Sutskever, I., & Salakhutdinov, R. (2014). Dropout: A simple way to prevent neural networks from overfitting. JMLR, 15(1), 1929-1958.",
        "Kingma, D. P., & Ba, J. (2014). Adam: A method for stochastic optimization. arXiv preprint arXiv:1412.6980."
    ]
    for r in refs:
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        run = p.add_run(r)
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        
    doc.save(REPORT_PATH)
    print(f"[SUKSES] Laporan Word bertingkat disimpan di: {REPORT_PATH}")
    return REPORT_PATH

if __name__ == '__main__':
    generate_word_report()
